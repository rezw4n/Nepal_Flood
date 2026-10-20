"""Build 'Nepal Flood Document.docx': the study as a conventional research article (numbered sections, methods
before results, APA author-year references), filled from the same computed results as the journal manuscript.

    python build_document.py            build from the cached reference metadata
    python build_document.py --fetch    first fetch missing author lists from Crossref (network)
"""
from pathlib import Path
import json
import re
import sys
import time
import pandas as pd
import requests
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_manuscript_v2 import facts, ROOT, DER, TAB, PAPER        # noqa: E402
from build_si import t_sources                               # noqa: E402
import export_v2 as ex                                       # noqa: E402

TITLE = "What height is safe? An analysis of building destruction along the 207-km path of Nepal's 2026 flood"
TEMPLATE = PAPER/'templates'/'document_template.md'
OUT_MD, OUT_DOCX = PAPER/'Nepal Flood Document.md', PAPER/'Nepal Flood Document.docx'
CROSSREF = ROOT/'sources_ext'/'crossref_authors.json'

# references cited only in this document
EXTRA = [
    dict(key='Gerl2016', authors='Gerl, T.; Kreibich, H.; Franco, G.; Marechal, D.; Schröter, K.', year=2016,
         title='A review of flood loss models as basis for harmonization and benchmarking', journal='PLoS ONE',
         volume='11(7)', pages='e0159791', type='journal-article', doi='10.1371/journal.pone.0159791', url=''),
    dict(key='Romali2025', authors='Romali, N. S.; Sulong, S.; Kawasaki, A.', year=2025,
         title='A systematic review of flood damage assessment: insight for the data-scarce regions',
         journal='Water Resources Management', volume='39(10)', pages='4707-4734', type='journal-article',
         doi='10.1007/s11269-025-04265-9', url=''),
    dict(key='LiL2023', authors='Li, L.; Chang-Richards, A.; Boston, M.; Elwood, K.; Molina Hutt, C.', year=2023,
         title='Post-disaster functional recovery of the built environment: A systematic review and directions for future research',
         journal='International Journal of Disaster Risk Reduction', volume='95', pages='103899', type='journal-article',
         doi='10.1016/j.ijdrr.2023.103899', url=''),
    dict(key='WHO2026', authors='World Health Organization', year=2026, title='Nepal: Flash floods and health response',
         journal='World Health Organization', volume='', pages='', type='web',
         doi='', url='https://www.who.int/nepal/emergencies/2026-rasuwa-flash-floods'),
    dict(key='Dietrich2026ISPRS', authors='Dietrich, O. et al. (8 authors)', year=2026,
         title='The potential of Copernicus satellites for disaster response: Retrieving building damage from Sentinel-1 and Sentinel-2',
         journal='ISPRS Annals of the Photogrammetry, Remote Sensing and Spatial Information Sciences', volume='XI-2-2026',
         pages='831-838', type='journal-article', doi='10.5194/isprs-annals-XI-2-2026-831-2026', url=''),
]
SHORT = {'EMSR927': 'Copernicus EMS', 'CEMSdash2026': 'Copernicus EMS', 'UNOSAT2026': 'UNOSAT',
         'RDNA2026': 'NDRRMA & NPC', 'WHO2026': 'WHO'}
FULL = {'RDNA2026': 'National Disaster Risk Reduction and Management Authority (NDRRMA), & National Planning Commission (NPC), Government of Nepal',
        'UNOSAT2026': 'UNOSAT (UNITAR)', 'WHO2026': 'World Health Organization (WHO)',
        'EMSR927': 'Copernicus Emergency Management Service (Copernicus EMS)',
        'CEMSdash2026': 'Copernicus Emergency Management Service (Copernicus EMS)'}
PERSON = re.compile(r"^[^\W\d][\w'’\-]*(?: [^\W\d][\w'’\-]*)*, (?:[A-Z][a-z]?\.\s?-?)+$")


# ------------------------------------------------------------------------------------------- reference metadata
def load_refs():
    refs = {r['key']: r for r in json.loads((PAPER/'references.json').read_text(encoding='utf-8'))}
    for r in EXTRA:
        refs[r['key']] = r
    # full author list of the HiRISK report (8 authors)
    refs['HiRISK2026'] = dict(refs['HiRISK2026'], authors='Liu, Q.; Khatri, A.; Lord, A.; Ghouli, M.; Thapa, N.; '
                                                            'Steiner, J.; Immerzeel, W. W.; Allen, S.')
    return refs


def fetch_crossref(refs, keys):
    """Cache full author lists, issue and pages from Crossref for cited DOIs (DataCite DOIs are skipped)."""
    cache = json.loads(CROSSREF.read_text(encoding='utf-8')) if CROSSREF.exists() else {}
    for k in sorted(keys):
        doi = (refs[k].get('doi') or '').strip()
        if not doi or doi in cache:
            continue
        try:
            r = requests.get(f'https://api.crossref.org/works/{doi}', timeout=30,
                             headers={'User-Agent': 'nepal-flood-paper-builder/1.0'})
            if r.status_code != 200:
                cache[doi] = {'status': r.status_code}
                continue
            m = r.json()['message']
            cache[doi] = {'status': 200,
                          'authors': [[a.get('family', ''), a.get('given', ''), a.get('name', '')] for a in m.get('author', [])],
                          'issue': m.get('issue', ''), 'volume': m.get('volume', ''), 'page': m.get('page', ''),
                          'container': (m.get('container-title') or [''])[0]}
        except requests.RequestException as e:
            cache[doi] = {'status': str(e)[:80]}
        time.sleep(0.2)
    CROSSREF.write_text(json.dumps(cache, indent=1, ensure_ascii=False), encoding='utf-8')
    print('crossref cache:', sum(v.get('status') == 200 for v in cache.values()), 'of', len(cache), 'DOIs resolved')


def surname_or_org(p):
    p = p.strip()
    m = re.search(r'\(([A-Z][A-Z0-9&\-]{1,10})\)', p)                 # 'Team (HOT)' -> 'HOT'
    if m:
        return m.group(1)
    return p.split(',')[0].strip() if PERSON.match(p) else p.rstrip('.')


def people(a):
    """Split an author string into individual names (persons as 'Surname, I.' or organisations)."""
    a = re.sub(r'\s*\(\d+ authors\)', '', a).strip()
    etal = bool(re.search(r'et al\.?$', a))
    a = re.sub(r',?\s*et al\.?$', '', a).strip()
    if ';' in a:
        names = [x.strip() for x in a.split(';') if x.strip()]
    elif '&' in a and re.search(r', [A-Z]\.', a):
        names = [x.strip(' ,') for x in re.findall(r"[^,&]+?, (?:[A-Z][a-z]?\.\s?-?)+", a)]
    else:
        names = [a]
    return names, etal


def first_author(r):
    """(surname, initials) of a personal first author, or (None, None) for an organisation."""
    names, _ = people(r['authors'])
    n = names[0].strip()
    if not PERSON.match(n):
        return None, None
    sur, ini = n.split(',', 1)
    return sur.strip(), ini.strip()


def short_name(k, r, initials_for=frozenset()):
    if k in SHORT:
        return SHORT[k]
    names, etal = people(r['authors'])
    s = [surname_or_org(n) for n in names]
    sur, ini = first_author(r)
    if sur in initials_for:                    # APA: different first authors sharing a surname are given initials
        s[0] = f'{ini} {sur}'
    if etal or len(s) > 2:
        return f'{s[0]} et al.'
    return s[0] if len(s) == 1 else f'{s[0]} & {s[1]}'


def initials(given):
    out = []
    for part in given.replace('.', '. ').split():
        sub = [x for x in part.split('-') if x.strip('.')]
        if sub:
            out.append('-'.join(x.strip('.')[0].upper() + '.' for x in sub))
    return ' '.join(out)


def apa_author_list(k, r, cr):
    if k in FULL:
        return FULL[k]
    names, etal = people(r['authors'])
    cra = (cr or {}).get('authors') or []
    # curated full author lists are kept; Crossref only completes 'et al.' lists whose first surname agrees
    if etal and cra and cra[0][0] and cra[0][0].lower() == surname_or_org(names[0]).lower():
        names = [(f'{fam}, {initials(giv)}' if fam else nm).strip(', ') for fam, giv, nm in cra]
        etal = False
    if etal:
        return f'{names[0]}, et al.'
    if len(names) == 1:
        return names[0]
    if len(names) > 20:
        return ', '.join(names[:19]) + ', . . . ' + names[-1]
    return ', '.join(names[:-1]) + ', & ' + names[-1]


def apa_entry(k, r, cr, suffix):
    au = apa_author_list(k, r, cr).rstrip('.')
    head = f"{au}. ({r['year']}{suffix})."
    title = r['title'].rstrip('.')
    vol, issue = '', ''
    m = re.match(r'^([^()]+)\((.+)\)$', str(r.get('volume') or ''))
    vol, issue = (m.group(1), m.group(2)) if m else (str(r.get('volume') or ''), '')
    if cr and cr.get('status') == 200:
        vol = vol or cr.get('volume', ''); issue = issue or cr.get('issue', '')
    pages = (r.get('pages') or (cr or {}).get('page') or '').replace('-', '–')
    if r.get('type') == 'book-chapter' and r.get('book_apa'):
        body = f"{title}. In {r['editors_apa']} (Eds.), *{r['book_apa']}* (pp. {pages}). {r['publisher']}."
    elif r.get('type') == 'journal-article' and r.get('journal'):
        src = f"*{r['journal']}*" + (f", *{vol}*" if vol else '') + (f"({issue})" if issue else '') + (f", {pages}" if pages else '')
        body = f'{title}. {src}.'
    else:
        body = f"*{title}*. {r.get('journal') or ''}."
    body = body.replace('..', '.').replace(' .', '.')                   # tidy the title/source only, not the author list
    link = f"https://doi.org/{r['doi']}" if r.get('doi') else (r.get('url') or '')
    target = link.replace('(', '%28').replace(')', '%29')               # DOIs such as S0734-189X(84)80011-0
    return f'{head} {body}' + (f' [{link}]({target})' if link else '')


# ------------------------------------------------------------------------------------------- citations
def cite_apa(text, refs, cache):
    keys = sorted(set(re.findall(r'@([A-Za-z][\w]*)', ' '.join(re.findall(r'\[@[^\]]+\]|\{@\w+\}', text)))))
    missing = [k for k in keys if k not in refs]
    if missing:
        raise KeyError(f'unknown reference keys: {missing}')
    persons = {}
    for k in keys:
        sur, ini = first_author(refs[k])
        if sur and k not in SHORT:
            persons.setdefault(sur, set()).add(ini)
    shared = frozenset(sur for sur, inis in persons.items() if len(inis) > 1)
    short = {k: short_name(k, refs[k], shared) for k in keys}
    groups = {}
    for k in keys:
        groups.setdefault((short[k], str(refs[k]['year'])), []).append(k)
    suffix = {}
    for (_, _), ks in groups.items():
        ks = sorted(ks, key=lambda k: refs[k]['title'].lower())
        for i, k in enumerate(ks):
            suffix[k] = 'abcdefghij'[i] if len(ks) > 1 else ''

    def paren(m):
        ks = [x.strip().lstrip('@') for x in m.group(1).split(';')]
        surname = lambda a: re.sub(r'^(?:[A-Z][a-z]?\.[\s-]?)+', '', a).lower()   # 'R. Li' sorts under L
        items = sorted({(short[k], f"{refs[k]['year']}{suffix[k]}") for k in ks}, key=lambda x: (surname(x[0]), x[1]))
        merged = {}
        for a, y in items:                     # one author, several works: 'Copernicus EMS, 2026a, 2026b'
            merged.setdefault(a, []).append(y)
        return '(' + '; '.join(f"{a}, {', '.join(ys)}" for a, ys in merged.items()) + ')'

    text = re.sub(r'\[(@[^\]]+)\]', paren, text)
    text = re.sub(r'\{@(\w+)\}', lambda m: f"{short[m.group(1)]} ({refs[m.group(1)]['year']}{suffix[m.group(1)]})", text)
    entries = []
    for k in keys:
        cr = cache.get((refs[k].get('doi') or '').strip())
        entries.append((apa_author_list(k, refs[k], cr).lower(), str(refs[k]['year']) + suffix[k],
                        apa_entry(k, refs[k], cr, suffix[k])))
    reflist = '\n\n'.join(e[2] for e in sorted(entries))
    return text, reflist, keys


# ------------------------------------------------------------------------------------------- content
def data_table():
    md = t_sources(json.loads((ROOT/'sources_ext'/'s2'/'s2_log.json').read_text()))
    out = []
    for i, line in enumerate(md.splitlines()):
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        if i > 1 and cells[0].startswith('Not used'):
            continue
        out.append('| ' + ' | '.join(cells[:-1]) + ' |')
    return '\n'.join(out)


def document_facts():
    F = facts()
    prof = pd.read_csv(DER/'flowpath_profile.csv')
    z = prof.set_index('chainage_m').z_channel
    zr = float(z.iloc[(z.index.to_series() - 23460).abs().argmin()])
    tr = pd.read_csv(DER/'transects.csv')
    vw = lambda a, b: tr[(tr.chainage_km >= a) & (tr.chainage_km < b)].valley_width_20m.median()
    ez = pd.read_csv(TAB/'exposure_growth_zones.csv')
    F.update(TITLE=TITLE, z_source=f'{round(z.iloc[0], -1):,.0f}', z_devghat=f'{round(z.iloc[-1], -1):,.0f}',
             drop_border=f'{(z.iloc[0] - zr)//100*100:,.0f}', vw_gorge=f'{round(vw(21, 41), -1):.0f}',
             vw_valley=f'{round(vw(64, 114), -1):.0f}',
             pop_1km=f"{round(ez[ez.zone == 'corridor_d1km'].pop_2026.sum(), -3):,.0f}", TABLE_DATA=data_table())
    return F


def check_order(md, label, n):
    body = md.split('## References')[0]
    seen = []
    for m in re.finditer(rf'(?<!\*\*)(?<!Supplementary ){label} (\d+)', body):
        k = int(m.group(1))
        if k not in seen:
            seen.append(k)
    if seen != list(range(1, n + 1)):
        raise ValueError(f'{label} first mentions out of order: {seen}')


# ------------------------------------------------------------------------------------------- Word export
def to_docx(bs, path, title):
    doc = Document(); sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15
    for name, size, italic in [('Title', 16, False), ('Heading 1', 14, False), ('Heading 2', 12, False), ('Heading 3', 12, True)]:
        h = doc.styles[name]; h.font.name = 'Times New Roman'; h.font.size = Pt(size); h.font.bold = True
        h.font.italic = italic; h.font.color.rgb = RGBColor(0, 0, 0)
        h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman'); ex.clear_theme_fonts(h)
        h.paragraph_format.space_before = Pt(12 if name != 'Title' else 0); h.paragraph_format.space_after = Pt(6)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); fp._p.append(fld)
    # layout: centred author block (9 pt) and 'Abstract' heading, abstract indented 0.3 in on both sides,
    # justified text except keywords, end-matter declarations and references
    LEFT_SECTIONS = {'acknowledgements', 'funding', 'author contributions', 'competing interests', 'references'}
    in_refs = False
    section, prev, authors, front = '', None, '', False
    i = 0
    while i < len(bs):
        kind, v = bs[i]
        if kind == 'heading':
            level, s = v
            section = s.strip().lower()
            in_refs = section == 'references'
            if level == 1:
                p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p = doc.add_paragraph(style=f'Heading {min(level - 1, 3)}')
                if section == 'abstract':
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ex.add_runs(p, s)
            prev = ('title' if level == 1 else 'heading', section)
            i += 1
            continue
        elif kind == 'pagebreak':
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == 'paragraph':
            p = doc.add_paragraph()
            justify = section not in LEFT_SECTIONS and not v.startswith('**Keywords')
            front = bool(prev and prev[0] in ('title', 'front'))
            if front:                                                # author line, then affiliation lines
                if prev[0] == 'title':
                    authors = re.sub(r'\^[^^]*\^|[*]', '', v)
                ex.add_runs(p, v.replace('**', ''), size=9); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if prev[0] == 'front':
                    p.paragraph_format.space_after = Pt(1)
                justify = False
            elif prev == ('heading', 'abstract'):                    # abstract paragraph
                ex.add_runs(p, v); p.paragraph_format.left_indent = p.paragraph_format.right_indent = Inches(0.3)
            elif ex.is_caption(v):
                ex.add_runs(p, v, size=10); p.paragraph_format.space_after = Pt(12); p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.keep_together = True
                if i + 1 < len(bs) and bs[i + 1][0] == 'table':
                    p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(4)
            elif v.startswith('*Note'):
                ex.add_runs(p, v, size=9.5); p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(12)
            elif in_refs:
                ex.add_runs(p, v, size=11); pf = p.paragraph_format
                pf.left_indent = Inches(0.5); pf.first_line_indent = Inches(-0.5); pf.line_spacing = 1.0
            else:
                ex.add_runs(p, v)
            if justify:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind == 'image':
            f = ROOT/v[1]
            with PILImage.open(f) as im:
                w, h = im.size
            width = min(6.5, 7.0*w/h)                                 # leaves room for the caption on the page
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(f), width=Inches(width)); p.paragraph_format.keep_with_next = True
        elif kind == 'table':
            ncol = len(v[0]); size = ex.fit_table_size(v, 6.5)
            t = doc.add_table(rows=0, cols=ncol); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, row in enumerate(v):
                cells = t.add_row().cells
                for cell, s in zip(cells, row + ['']*(ncol - len(row))):
                    cp = cell.paragraphs[0]; cp.paragraph_format.space_after = Pt(1); cp.paragraph_format.line_spacing = 1.0
                    ex.add_runs(cp, s, size=size)
                    if j == 0:
                        cp.paragraph_format.keep_with_next = True
                        for rr in cp.runs:
                            rr.bold = True
                        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'E7EEF3')
                        cell._tc.get_or_add_tcPr().append(shd)
                trp = t.rows[-1]._tr.get_or_add_trPr(); trp.append(OxmlElement('w:cantSplit'))
                if j == 0:
                    trp.append(OxmlElement('w:tblHeader'))
            ex.set_table_widths(t, v, 6.5, size_pt=size)
        prev = ('front', section) if kind == 'paragraph' and front else (kind, section)
        i += 1
    cp = doc.core_properties
    cp.title = title; cp.subject = '26 August 2026 Bhote Koshi–Trishuli flood, Nepal'
    cp.author = '; '.join(a.strip() for a in authors.split(',')) if authors and 'supplied' not in authors else 'Author list pending'
    doc.save(path)


def build(fetch=False):
    refs = load_refs()
    F = document_facts()
    t = TEMPLATE.read_text(encoding='utf-8')
    missing = sorted(set(re.findall(r'\{\{(\w+)\}\}', t)) - set(F) - {'REFERENCES'})
    if missing:
        raise KeyError(f'unresolved placeholders: {missing}')
    for k, v in F.items():
        t = t.replace('{{' + k + '}}', str(v))
    if fetch:
        keys = set(re.findall(r'@([A-Za-z]\w*)', ' '.join(re.findall(r'\[@[^\]]+\]|\{@\w+\}', t))))
        fetch_crossref(refs, keys)
    cache = json.loads(CROSSREF.read_text(encoding='utf-8')) if CROSSREF.exists() else {}
    t, reflist, keys = cite_apa(t, refs, cache)
    t = re.sub(r'(?<![\w/.])-(?=\d)', '−', t)                    # typographic minus (before the reference list)
    t = t.replace('{{REFERENCES}}', reflist)
    if re.search(r'\{\{\w+\}\}|\[@|\{@', t):
        raise ValueError('unresolved placeholders or citations remain')
    check_order(t, 'Figure', 8); check_order(t, 'Table', 3)
    OUT_MD.write_text(t, encoding='utf-8')
    to_docx(ex.blocks(t), OUT_DOCX, TITLE)
    words = len(t.split('## 1. Introduction')[1].split('## Data availability')[0].split())
    print(f'wrote {OUT_DOCX.name}: {len(keys)} references (APA), about {words:,} words from Introduction to Conclusion')


if __name__ == '__main__':
    build(fetch='--fetch' in sys.argv)
