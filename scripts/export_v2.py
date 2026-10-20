"""Export the manuscript and Supplementary Information (Markdown subset) to DOCX, PDF and self-contained HTML.

Markdown subset: '#'-headings, paragraphs, pipe tables, images ![alt](path), '\\pagebreak',
inline **bold**, *italic*, ^superscript^, ~subscript~ and [links](url).
"""
from pathlib import Path
import base64
import html
import re
import sys
from datetime import datetime, timezone
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

ROOT = Path(__file__).resolve().parents[1]
RUNNING_TITLE = "What height is safe? Nepal's 2026 flood"
TOKEN = re.compile(r'(\*\*.+?\*\*|\*[^*\s][^*]*?\*|\^[^^\s][^^]*?\^|~[^~\s][^~]*?~|\[[^\]]+\]\([^)]+\))')
HEAD = '#183E54'


def blocks(md):
    md = re.sub(r'<!--.*?-->', '', md, flags=re.S)
    lines = md.splitlines(); out = []; i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1; continue
        if s == r'\pagebreak':
            out.append(('pagebreak', None)); i += 1; continue
        if s.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', c) for c in row):
                    rows.append(row)
                i += 1
            out.append(('table', rows)); continue
        if s.startswith('!['):
            m = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', s); out.append(('image', (m[1], m[2]))); i += 1; continue
        if s.startswith('#'):
            m = re.match(r'(#+)\s+(.*)', s); out.append(('heading', (len(m[1]), m[2]))); i += 1; continue
        para = [s]; i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].lstrip().startswith(('#', '|', '![', r'\pagebreak')):
            para.append(lines[i].strip()); i += 1
        out.append(('paragraph', ' '.join(para)))
    return out


def tokens(s):
    """Yield (text, style dict) runs."""
    pos = 0
    for m in TOKEN.finditer(s):
        if m.start() > pos:
            yield s[pos:m.start()], {}
        t = m.group(0)
        if t.startswith('**'):
            for sub, st in tokens(t[2:-2]):
                yield sub, dict(st, bold=True)
        elif t.startswith('*'):
            for sub, st in tokens(t[1:-1]):
                yield sub, dict(st, italic=True)
        elif t.startswith('^'):
            yield t[1:-1], {'sup': True}
        elif t.startswith('~'):
            yield t[1:-1], {'sub': True}
        else:
            lm = re.match(r'\[([^\]]+)\]\(([^)]+)\)', t)
            yield lm[1], {'link': lm[2]}
        pos = m.end()
    if pos < len(s):
        yield s[pos:], {}


def rl_inline(s):
    out = []
    for t, st in tokens(s):
        x = html.escape(t)
        if st.get('link'):
            x = f'<a href="{html.escape(st["link"], quote=True)}" color="#156080">{x}</a>'
        if st.get('sup'):
            x = f'<super>{x}</super>'
        if st.get('sub'):
            x = f'<sub>{x}</sub>'
        if st.get('italic'):
            x = f'<i>{x}</i>'
        if st.get('bold'):
            x = f'<b>{x}</b>'
        out.append(x)
    return ''.join(out)


def html_inline(s):
    out = []
    for t, st in tokens(s):
        x = html.escape(t)
        if st.get('link'):
            x = f'<a href="{html.escape(st["link"], quote=True)}">{x}</a>'
        if st.get('sup'):
            x = f'<sup>{x}</sup>'
        if st.get('sub'):
            x = f'<sub>{x}</sub>'
        if st.get('italic'):
            x = f'<em>{x}</em>'
        if st.get('bold'):
            x = f'<strong>{x}</strong>'
        out.append(x)
    return ''.join(out)


def is_caption(s):
    return bool(re.match(r'^\*\*(Fig\.|Figure|Table|Supplementary (Fig\.|Figure|Table))\s*\S+', s))


# ------------------------------------------------------------------------------------------- Word
def add_runs(p, s, size=None, color=None):
    for t, st in tokens(s):
        if st.get('link'):
            rel = p.part.relate_to(st['link'], 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                                   is_external=True)
            a = OxmlElement('w:hyperlink'); a.set(qn('r:id'), rel)
            r = OxmlElement('w:r'); rp = OxmlElement('w:rPr'); c = OxmlElement('w:color'); c.set(qn('w:val'), '156080')
            rp.append(c)
            if size:                                   # links take the size of the surrounding text (w:sz follows w:color)
                for tag in ('w:sz', 'w:szCs'):
                    e = OxmlElement(tag); e.set(qn('w:val'), str(int(round(size*2)))); rp.append(e)
            r.append(rp); te = OxmlElement('w:t'); te.text = t; te.set(qn('xml:space'), 'preserve')
            r.append(te); a.append(r); p._p.append(a); continue
        run = p.add_run(t)
        run.bold = st.get('bold'); run.italic = st.get('italic')
        if st.get('sup'):                      # set only when needed: both share w:vertAlign, and None clears it
            run.font.superscript = True
        elif st.get('sub'):
            run.font.subscript = True
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def clear_theme_fonts(style):
    """Remove theme-font attributes, which override the explicit font in Word; drop the Title rule."""
    rf = style.element.get_or_add_rPr().find(qn('w:rFonts'))
    if rf is not None:
        for att in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            rf.attrib.pop(qn(att), None)
    ppr = style.element.find(qn('w:pPr'))
    if ppr is not None and ppr.find(qn('w:pBdr')) is not None:
        ppr.remove(ppr.find(qn('w:pBdr')))


def column_needs(rows, size_pt):
    """Minimum width of each column (its longest word in Times, bold in the header row, plus cell margins, in inches)
    and a weight for sharing the remaining width (length of the cell text)."""
    plain = [[re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', re.sub(r'[*^~]', '', c)) for c in r] for r in rows]
    ncol = max(len(r) for r in plain)
    minw, flex = [], []
    for k in range(ncol):
        col = [r[k] if k < len(r) else '' for r in plain]
        need = max(pdfmetrics.stringWidth(w, 'Times-Bold' if j == 0 else 'Times-Roman', size_pt)
                   for j, c in enumerate(col) for w in (c.split() or ['']))
        dlen = max([len(c) for c in col[1:]] or [0])
        minw.append(min(need/72, 1.6) + 0.16)
        flex.append(max(min(dlen, 40) - need/72/0.05, 1))
    return minw, flex


def fit_table_size(rows, total_in, sizes=(8.5, 8.0, 7.5, 7.0), spare_in=0.6):
    """Largest font size at which no word is broken and some width is left to share between columns."""
    for s in sizes:
        if sum(column_needs(rows, s)[0]) + spare_in <= total_in:
            return s
    return sizes[-1]


def set_table_widths(table, rows, total_in, size_pt=8.0):
    """Column widths from content: each column at least as wide as its longest word, the rest shared by text length."""
    minw, flex = column_needs(rows, size_pt)
    extra = total_in - sum(minw)
    widths = [m + max(extra, 0)*f/sum(flex) for m, f in zip(minw, flex)] if extra > 0 else [total_in*m/sum(minw) for m in minw]
    table.autofit = False                                   # writes <w:tblLayout w:type="fixed"/>
    for row in table.rows:
        for k, cell in enumerate(row.cells):
            cell.width = Inches(widths[k])
    for k, col in enumerate(table.columns):
        col.width = Inches(widths[k])


def to_word(bs, path, title, header):
    doc = Document(); sec = doc.sections[0]
    sec.page_width = Inches(8.2677); sec.page_height = Inches(11.6929)
    sec.top_margin = sec.bottom_margin = Inches(0.9); sec.left_margin = sec.right_margin = Inches(0.95)
    st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.5
    for name, size in [('Title', 18), ('Heading 1', 13.5), ('Heading 2', 12), ('Heading 3', 11)]:
        h = doc.styles[name]; h.font.name = 'Arial'; h.font.size = Pt(size); h.font.color.rgb = RGBColor.from_string(HEAD[1:])
        h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial'); clear_theme_fonts(h)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(header + '  |  '); r.font.size = Pt(8)
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); fp._p.append(fld)
    ln = OxmlElement('w:lnNumType'); ln.set(qn('w:countBy'), '1'); ln.set(qn('w:restart'), 'continuous')
    sec._sectPr.append(ln)
    i = 0
    while i < len(bs):
        kind, v = bs[i]
        if kind == 'heading':
            level, s = v
            hp = doc.add_heading(level=0 if level == 1 else min(level - 1, 3))
            add_runs(hp, s)
        elif kind == 'pagebreak':
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == 'paragraph':
            p = doc.add_paragraph()
            if is_caption(v):
                add_runs(p, v, size=9.5); p.paragraph_format.line_spacing = 1.15; p.paragraph_format.space_after = Pt(12)
            else:
                add_runs(p, v)
        elif kind == 'image':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ROOT/v[1]), width=Inches(6.3)); p.paragraph_format.keep_with_next = True
        elif kind == 'table':
            ncol = len(v[0])
            tw = 8.2677 - 2*0.95
            size = fit_table_size(v, tw, sizes=(8.0, 7.5, 7.0))
            t = doc.add_table(rows=0, cols=ncol); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, row in enumerate(v):
                cells = t.add_row().cells
                for cell, s in zip(cells, row + ['']*(ncol - len(row))):
                    cell.paragraphs[0].paragraph_format.space_after = Pt(1)
                    cell.paragraphs[0].paragraph_format.line_spacing = 1.0
                    add_runs(cell.paragraphs[0], s, size=size)
                    if j == 0:
                        cell.paragraphs[0].paragraph_format.keep_with_next = True
                        for rr in cell.paragraphs[0].runs:
                            rr.bold = True
                        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'E6EEF2')
                        cell._tc.get_or_add_tcPr().append(shd)
                tr = t.rows[-1]._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); tr.append(cs)
                if j == 0:
                    tr.append(OxmlElement('w:tblHeader'))
            set_table_widths(t, v, tw, size_pt=size)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        i += 1
    doc.core_properties.title = title
    doc.core_properties.author = authors_of(bs)
    doc.core_properties.created = doc.core_properties.modified = datetime.now(timezone.utc); doc.core_properties.revision = 1
    doc.save(path)


# ------------------------------------------------------------------------------------------- PDF
def register_fonts():
    fd = Path('C:/Windows/Fonts')
    for name, f in [('Body', 'times.ttf'), ('BodyBold', 'timesbd.ttf'), ('BodyItalic', 'timesi.ttf'),
                    ('BodyBI', 'timesbi.ttf'), ('Head', 'arial.ttf'), ('HeadBold', 'arialbd.ttf'),
                    ('HeadItalic', 'ariali.ttf'), ('HeadBI', 'arialbi.ttf')]:
        pdfmetrics.registerFont(TTFont(name, str(fd/f)))
    pdfmetrics.registerFontFamily('Body', normal='Body', bold='BodyBold', italic='BodyItalic', boldItalic='BodyBI')
    pdfmetrics.registerFontFamily('Head', normal='Head', bold='HeadBold', italic='HeadItalic', boldItalic='HeadBI')


def to_pdf(bs, path, title, header):
    register_fonts()
    sty = {
        'p': ParagraphStyle('p', fontName='Body', fontSize=10.3, leading=14.2, spaceAfter=6.5, alignment=TA_JUSTIFY),
        'cap': ParagraphStyle('cap', fontName='Head', fontSize=7.9, leading=10.4, spaceAfter=10, alignment=TA_LEFT),
        'capk': ParagraphStyle('capk', fontName='Head', fontSize=7.9, leading=10.4, spaceAfter=4, alignment=TA_LEFT,
                               keepWithNext=True),                  # table title stays with its table
        'h1': ParagraphStyle('h1', fontName='HeadBold', fontSize=17, leading=21, spaceAfter=10, textColor=colors.HexColor(HEAD)),
        'h2': ParagraphStyle('h2', fontName='HeadBold', fontSize=12.5, leading=16, spaceBefore=11, spaceAfter=6,
                             keepWithNext=True, textColor=colors.HexColor(HEAD)),
        'h3': ParagraphStyle('h3', fontName='HeadBold', fontSize=10.6, leading=13.5, spaceBefore=8, spaceAfter=4,
                             keepWithNext=True, textColor=colors.HexColor(HEAD)),
        'h4': ParagraphStyle('h4', fontName='HeadBold', fontSize=9.8, leading=12.5, spaceBefore=6, spaceAfter=3,
                             keepWithNext=True, textColor=colors.HexColor('#1a1a1a')),
        'cell': ParagraphStyle('cell', fontName='Head', fontSize=7.0, leading=8.8),
        'cellh': ParagraphStyle('cellh', fontName='HeadBold', fontSize=7.0, leading=8.8, textColor=colors.HexColor('#1a1a1a')),
        'ref': ParagraphStyle('ref', fontName='Body', fontSize=8.8, leading=11.2, spaceAfter=3.2, alignment=TA_LEFT),
    }
    width = A4[0] - 2*56
    flows = []; i = 0; in_refs = False; pending = None
    while i < len(bs):
        kind, v = bs[i]
        if kind == 'heading':
            level, s = v
            in_refs = s.strip().lower() in ('references', 'supplementary references')
            flows.append(Paragraph(rl_inline(s), sty[f'h{min(level, 4)}']))
        elif kind == 'pagebreak':
            flows.append(PageBreak())
        elif kind == 'paragraph':
            if is_caption(v) and i + 1 < len(bs) and bs[i + 1][0] == 'table':
                pending = Paragraph(rl_inline(v), sty['capk'])     # placed with its table below
            else:
                flows.append(Paragraph(rl_inline(v), sty['cap' if is_caption(v) else ('ref' if in_refs else 'p')]))
        elif kind == 'image':
            p = ROOT/v[1]
            with PILImage.open(p) as im:
                w, h = im.size
            hmax = A4[1] - 210
            ww = min(width, hmax*w/h)
            together = [Spacer(1, 4), Image(str(p), width=ww, height=ww*h/w), Spacer(1, 4)]
            if i + 1 < len(bs) and bs[i + 1][0] == 'paragraph' and is_caption(bs[i + 1][1]):
                i += 1; together.append(Paragraph(rl_inline(bs[i][1]), sty['cap']))
            flows.append(KeepTogether(together))
        elif kind == 'table':
            ncol = len(v[0])
            data = [[Paragraph(rl_inline(c), sty['cellh' if j == 0 else 'cell']) for c in row + ['']*(ncol - len(row))]
                    for j, row in enumerate(v)]
            plain = [[re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', re.sub(r'[*^~]', '', c)) for c in row + ['']*(ncol - len(row))]
                     for row in v]
            weights, minw = [], []
            for k in range(ncol):
                dlen = max([len(r[k]) for r in plain[1:]] or [0])
                weights.append(max(min(dlen, 28), 6) + 2)
                need = max(pdfmetrics.stringWidth(w, 'HeadBold' if j == 0 else 'Head', 7.0)
                           for j, r in enumerate(plain) for w in (r[k].split() or ['']))
                minw.append(min(need, 110) + 8)                 # longest word plus cell padding (points)
            if sum(minw) < width:                               # reserve the longest word, share the rest by content
                flex = [max(w - m/3.5, 1) for w, m in zip(weights, minw)]
                widths = [m + (width - sum(minw))*f/sum(flex) for m, f in zip(minw, flex)]
            else:
                widths = [width*w/sum(weights) for w in weights]
            t = Table(data, colWidths=widths, repeatRows=1, hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e6eef2')),
                                   ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                                   ('LINEABOVE', (0, 0), (-1, 0), 0.8, colors.HexColor('#1a1a1a')),
                                   ('LINEBELOW', (0, 0), (-1, 0), 0.5, colors.HexColor('#1a1a1a')),
                                   ('LINEBELOW', (0, -1), (-1, -1), 0.8, colors.HexColor('#1a1a1a')),
                                   ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f6f6f3')]),
                                   ('LEFTPADDING', (0, 0), (-1, -1), 3.5), ('RIGHTPADDING', (0, 0), (-1, -1), 3.5),
                                   ('TOPPADDING', (0, 0), (-1, -1), 2.8), ('BOTTOMPADDING', (0, 0), (-1, -1), 2.8)]))
            note = []
            if i + 1 < len(bs) and bs[i + 1][0] == 'paragraph' and bs[i + 1][1].startswith('*Note'):
                i += 1; note = [Paragraph(rl_inline(bs[i][1]), sty['cap'])]
            head = [pending] if pending else []
            pending = None
            if len(v) < 12:
                flows.append(KeepTogether(head + [Spacer(1, 3), t, Spacer(1, 4)] + note))
            else:      # tall tables may split (header row repeats); the title is kept with the first rows
                flows.extend(head + [t, Spacer(1, 4)] + note)
        i += 1

    def page(c, d):
        c.saveState(); c.setFont('Head', 7.2); c.setFillColor(colors.HexColor('#6b6b66'))
        c.drawString(56, 30, header); c.drawRightString(A4[0] - 56, 30, str(d.page)); c.restoreState()
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=56, rightMargin=56, topMargin=50, bottomMargin=48,
                            title=title, author=authors_of(bs))
    doc.build(flows, onFirstPage=page, onLaterPages=page)


# ------------------------------------------------------------------------------------------- HTML
CSS = """
:root{--ink:#1a1a1a;--ink2:#52514e;--muted:#898781;--line:#e1e0d9;--accent:#183e54;--surface:#ffffff;--band:#f6f6f3;--link:#156080}
@media (prefers-color-scheme: dark){:root:not([data-theme="light"]){--ink:#f0efec;--ink2:#c3c2b7;--muted:#898781;--line:#383835;--accent:#9cc3dc;--surface:#1a1a19;--band:#232321;--link:#8fc1e0}}
:root[data-theme="dark"]{--ink:#f0efec;--ink2:#c3c2b7;--muted:#898781;--line:#383835;--accent:#9cc3dc;--surface:#1a1a19;--band:#232321;--link:#8fc1e0}
*{box-sizing:border-box}html{background:var(--surface)}
body{max-width:920px;margin:0 auto;padding:40px 16px 80px;background:var(--surface);color:var(--ink);font:17px/1.62 Georgia,'Times New Roman',serif}
h1,h2,h3,h4{font-family:Arial,Helvetica,sans-serif;color:var(--accent);line-height:1.25}
h1{font-size:30px;margin:0 0 14px}h2{font-size:22px;margin-top:40px;border-top:1px solid var(--line);padding-top:16px}
h3{font-size:18px;margin-top:28px}h4{font-size:16px;color:var(--ink)}
a{color:var(--link)}sup{font-size:.68em;line-height:0}
figure{margin:26px 0}figure img{width:100%;height:auto;background:#fff;border:1px solid var(--line)}
figcaption,.note{font:14px/1.5 Arial,Helvetica,sans-serif;color:var(--ink2)}
.tablewrap{overflow-x:auto;margin:18px 0}
table{border-collapse:collapse;width:100%;font:13px/1.4 Arial,Helvetica,sans-serif}
th,td{padding:6px 8px;text-align:left;border-bottom:1px solid var(--line);vertical-align:top}
th{border-top:2px solid var(--ink);border-bottom:1px solid var(--ink)}tr:nth-child(even) td{background:var(--band)}
.ref{font-size:14px;line-height:1.45}
@media (max-width:600px){body{font-size:16px}h1{font-size:24px}}
"""


def to_html(bs, path, title):
    out = [f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
           f'<title>{html.escape(title)}</title><style>{CSS}</style></head><body>']
    i = 0; in_refs = False
    while i < len(bs):
        kind, v = bs[i]
        if kind == 'heading':
            level, s = v
            in_refs = s.strip().lower() in ('references', 'supplementary references')
            out.append(f'<h{level}>{html_inline(s)}</h{level}>')
        elif kind == 'paragraph':
            cls = ' class="note"' if is_caption(v) or v.startswith('*Note') else (' class="ref"' if in_refs else '')
            out.append(f'<p{cls}>{html_inline(v)}</p>')
        elif kind == 'image':
            data = base64.b64encode((ROOT/v[1]).read_bytes()).decode()
            cap = ''
            if i + 1 < len(bs) and bs[i + 1][0] == 'paragraph' and is_caption(bs[i + 1][1]):
                i += 1; cap = f'<figcaption>{html_inline(bs[i][1])}</figcaption>'
            out.append(f'<figure><img alt="{html.escape(v[0])}" src="data:image/png;base64,{data}">{cap}</figure>')
        elif kind == 'table':
            rows = ''.join('<tr>' + ''.join(f'<{"th" if j == 0 else "td"}>{html_inline(c)}</{"th" if j == 0 else "td"}>'
                                            for c in row) + '</tr>' for j, row in enumerate(v))
            out.append(f'<div class="tablewrap"><table>{rows}</table></div>')
        i += 1
    out.append('</body></html>')
    path.write_text('\n'.join(out), encoding='utf-8')


def authors_of(bs):
    """Author names from the first paragraph with affiliation markers (^1^), for the file properties."""
    for kind, v in bs:
        if kind == 'paragraph' and '^1^' in v:
            return '; '.join(a.strip() for a in re.sub(r'\^[^^]*\^|[*]', '', v).split(',') if a.strip())
    return ''


def export(md_path, stem, header):
    bs = blocks(md_path.read_text(encoding='utf-8'))
    title = re.sub(r'[*^~]', '', bs[0][1][1]) if bs and bs[0][0] == 'heading' else stem
    to_word(bs, md_path.with_suffix('.docx'), title, header)
    to_pdf(bs, md_path.with_suffix('.pdf'), title, header)
    to_html(bs, md_path.with_suffix('.html'), title)
    print('exported', stem)


if __name__ == '__main__':
    export(ROOT/'paper'/'Nepal_Flood_Damage_Assessment.md', 'Nepal_Flood_Damage_Assessment', RUNNING_TITLE)
    if (ROOT/'paper'/'Supplementary_Information.md').exists():
        export(ROOT/'paper'/'Supplementary_Information.md', 'Supplementary_Information', 'Supplementary Information')
